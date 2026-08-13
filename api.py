"""开源版后端：数据接收 + 视频列表/详情 + 个人分析 + 定时清理 + 导出。"""
import csv
import io
import json
import os
import tempfile
import threading
import time as time_module
from contextlib import asynccontextmanager
from datetime import datetime
from typing import Optional

import pymysql
from fastapi import FastAPI, Query, HTTPException, Header, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, RedirectResponse, StreamingResponse
from pydantic import BaseModel

import analyzer
import cleanup_service
import export_service
import extension_receiver
from time_filter import build_publish_filter

MYSQL_HOST = 'localhost'
MYSQL_PORT = 3307
MYSQL_USER = 'root'
MYSQL_PASSWORD = ''
MYSQL_DB = 'douyin_spider'

try:
    from local_config import (
        MYSQL_HOST as _H, MYSQL_PORT as _P, MYSQL_USER as _U,
        MYSQL_PASSWORD as _PW, MYSQL_DB as _DB,
        EXTENSION_API_TOKEN as _TOK, ALLOWED_AUTHOR_IDS as _IDS, CLEANUP_STORAGE as _STORE,
    )
    MYSQL_HOST, MYSQL_PORT, MYSQL_USER, MYSQL_PASSWORD, MYSQL_DB = _H, _P, _U, _PW, _DB
    EXTENSION_API_TOKEN, ALLOWED_AUTHOR_IDS, CLEANUP_STORAGE = _TOK, _IDS, _STORE
except ImportError:
    EXTENSION_API_TOKEN = ''
    ALLOWED_AUTHOR_IDS = []
    CLEANUP_STORAGE = 'json'

ALLOWED_ORIGINS = [
    'http://127.0.0.1:8002',
    'http://localhost:8002',
    'http://localhost:5173',
]
CLEANUP_CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'cleanup_config.json')


def _read_cleanup_config() -> dict:
    if CLEANUP_STORAGE == 'json':
        return cleanup_service.read_cleanup_config(CLEANUP_CONFIG_PATH)
    raise RuntimeError('开源版仅支持 CLEANUP_STORAGE=json')


def _write_cleanup_config(cfg: dict) -> None:
    if CLEANUP_STORAGE == 'json':
        cleanup_service.write_cleanup_config(CLEANUP_CONFIG_PATH, cfg)
        return
    raise RuntimeError('开源版仅支持 CLEANUP_STORAGE=json')


def _cleanup_once() -> None:
    cfg = _read_cleanup_config()
    enabled = bool(cfg['enabled'])
    last_raw = cfg['last_clean_time']
    last = None
    if last_raw:
        try:
            last = datetime.fromisoformat(last_raw)
        except ValueError:
            last = None
    if not cleanup_service.should_run_cleanup(enabled, last, datetime.now()):
        return
    batch_size = int(cfg['batch_size'])
    authors = list(cfg['authors'])

    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute('SELECT video_id, author_id, update_time FROM video_info')
            light_rows = cursor.fetchall()
    finally:
        db_close(db)
    if not light_rows:
        return

    ids = cleanup_service.select_stale_ids_per_author(
        light_rows, batch_size=batch_size, author_ids=authors or None,
    )
    if not ids:
        print('定时清理跳过：没有满足条件的待删数据')
        return

    db = get_db()
    try:
        with db.cursor() as cursor:
            placeholders = ', '.join(['%s'] * len(ids))
            cursor.execute(
                f'SELECT * FROM video_info WHERE video_id IN ({placeholders})',
                tuple(ids),
            )
            delete_rows = cursor.fetchall()
    finally:
        db_close(db)

    backup_dir = cleanup_service.CLEANUP_BACKUP_DIR
    os.makedirs(backup_dir, exist_ok=True)
    backup_path = os.path.join(
        backup_dir,
        'cleanup_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '.csv',
    )
    with open(backup_path, 'w', encoding='utf-8', newline='') as f:
        f.write(cleanup_service.build_backup_csv(delete_rows))

    db = get_db()
    try:
        with db.cursor() as cursor:
            placeholders = ', '.join(['%s'] * len(ids))
            cursor.execute(
                f'DELETE FROM video_info WHERE video_id IN ({placeholders})',
                tuple(ids),
            )
            db.commit()
    finally:
        db_close(db)

    cfg['last_clean_time'] = datetime.now().isoformat(timespec='seconds')
    _write_cleanup_config(cfg)
    print(f'定时清理完成：删除 {len(ids)} 条，备份 {backup_path}')


def _cleanup_loop() -> None:
    while True:
        try:
            _cleanup_once()
        except Exception as e:  # noqa: BLE001
            print(f'定时清理异常：{e}')
        time_module.sleep(24 * 3600)


@asynccontextmanager
async def lifespan(_app):
    threading.Thread(target=_cleanup_loop, daemon=True).start()
    yield


app = FastAPI(title='抖音创作者数据分析器', version='1.0.0', lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=['*'],
    allow_headers=['*'],
)


def verify_write_guard(
    origin: Optional[str] = Header(default=None),
    x_api_token: Optional[str] = Header(default=None, alias='X-API-Token'),
) -> None:
    allowed, status_code, reason = extension_receiver.evaluate_write_guard(
        origin, x_api_token, EXTENSION_API_TOKEN, ALLOWED_ORIGINS,
    )
    if not allowed:
        raise HTTPException(status_code=status_code, detail=reason)


def get_db():
    try:
        return pymysql.connect(
            host=MYSQL_HOST,
            port=MYSQL_PORT,
            user=MYSQL_USER,
            password=MYSQL_PASSWORD,
            database=MYSQL_DB,
            charset='utf8mb4',
            cursorclass=pymysql.cursors.DictCursor,
            connect_timeout=5,
        )
    except pymysql.Error as e:
        raise HTTPException(status_code=503, detail=f'MySQL 连接失败: {e}')


def db_close(db):
    try:
        db.close()
    except Exception:
        pass


def apply_publish_filter(start_date: str, end_date: str):
    try:
        return build_publish_filter(start_date or None, end_date or None)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


def _check_export_total(total: int) -> None:
    if total > export_service.EXPORT_MAX_ROWS:
        raise HTTPException(status_code=400, detail=f'数据量过大（{total} 条），请缩小筛选范围后导出')


class VideoItem(BaseModel):
    video_id: str
    video_title: Optional[str] = None
    video_desc: Optional[str] = None
    author_name: Optional[str] = None
    author_id: Optional[str] = None
    publish_time: Optional[datetime] = None
    like_count: Optional[int] = None
    comment_count: Optional[int] = None
    share_count: Optional[int] = None
    collect_count: Optional[int] = None
    play_count: Optional[int] = None
    video_url: Optional[str] = None
    cover_url: Optional[str] = None
    crawl_time: Optional[datetime] = None
    update_time: Optional[datetime] = None


class PaginatedResponse(BaseModel):
    total: int
    page: int
    page_size: int
    total_pages: int
    data: list[VideoItem]


class ExtensionVideosRequest(BaseModel):
    source_url: str
    videos: list[dict]


class CleanupToggleRequest(BaseModel):
    enabled: bool


class CleanupSettingsRequest(BaseModel):
    batch_size: int = 200
    authors: list[str] = []


@app.get('/api/videos', response_model=PaginatedResponse)
def list_videos(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: str = Query('', description='搜索视频标题/作者/ID'),
    sort_by: str = Query('crawl_time', description='排序字段'),
    order: str = Query('desc', pattern='^(asc|desc)$'),
    start_date: str = Query('', description='发布时间起始（YYYY-MM-DD）'),
    end_date: str = Query('', description='发布时间结束（YYYY-MM-DD）'),
):
    allowed_sort = {
        'video_id', 'video_title', 'author_name', 'publish_time',
        'like_count', 'comment_count', 'share_count', 'play_count', 'collect_count',
        'crawl_time', 'update_time',
    }
    if sort_by not in allowed_sort:
        sort_by = 'crawl_time'
    order_clause = 'DESC' if order == 'desc' else 'ASC'
    offset = (page - 1) * page_size
    publish_clause, publish_params = apply_publish_filter(start_date, end_date)
    author_clause, author_params = extension_receiver.build_author_filter(ALLOWED_AUTHOR_IDS)
    where_parts = []
    params = []
    if search:
        where_parts.append('(video_id LIKE %s OR video_title LIKE %s OR author_name LIKE %s)')
        params.extend([f'%{search}%'] * 3)
    if publish_clause:
        where_parts.append(publish_clause)
        params.extend(publish_params)
    if author_clause:
        where_parts.append(author_clause)
        params.extend(author_params)
    where_sql = ('WHERE ' + ' AND '.join(where_parts)) if where_parts else ''

    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute(f'SELECT COUNT(*) AS total FROM video_info {where_sql}', tuple(params))
            total = cursor.fetchone()['total']
            cursor.execute(
                f'SELECT * FROM video_info {where_sql} ORDER BY {sort_by} {order_clause} LIMIT %s OFFSET %s',
                tuple(params + [page_size, offset]),
            )
            rows = cursor.fetchall()
        return PaginatedResponse(
            total=total,
            page=page,
            page_size=page_size,
            total_pages=(total + page_size - 1) // page_size,
            data=[VideoItem(**row) for row in rows],
        )
    finally:
        db_close(db)


@app.get('/api/videos/{video_id}', response_model=VideoItem)
def get_video(video_id: str):
    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute('SELECT * FROM video_info WHERE video_id = %s', (video_id,))
            row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail='视频不存在')
        return VideoItem(**row)
    finally:
        db_close(db)


@app.get('/api/analyze/authors')
def analyze_authors():
    db = get_db()
    try:
        with db.cursor() as cursor:
            author_clause, author_params = extension_receiver.build_author_filter(ALLOWED_AUTHOR_IDS)
            where_sql = 'author_id IS NOT NULL AND author_id <> \'\''
            params = []
            if author_clause:
                where_sql += ' AND ' + author_clause
                params.extend(author_params)
            cursor.execute(f"""
                SELECT author_id, author_name, COUNT(*) AS count
                FROM video_info
                WHERE {where_sql}
                GROUP BY author_id, author_name
                ORDER BY count DESC
            """, tuple(params))
            rows = cursor.fetchall()
    finally:
        db_close(db)
    return {'authors': rows}


@app.get('/api/analyze/personal')
def analyze_personal(
    author_id: str = Query(..., description='作者 uid'),
    sort_by: str = Query('likes', description='Top 视频排序维度'),
    start_date: str = Query('', description='发布时间起始（YYYY-MM-DD）'),
    end_date: str = Query('', description='发布时间结束（YYYY-MM-DD）'),
):
    if sort_by not in ('likes', 'plays', 'comments', 'shares', 'collects', 'engagement'):
        raise HTTPException(status_code=400, detail='sort_by 必须是 likes/plays/comments/shares/collects/engagement')
    db = get_db()
    try:
        with db.cursor() as cursor:
            publish_clause, publish_params = apply_publish_filter(start_date, end_date)
            author_clause, author_params = extension_receiver.build_author_filter(ALLOWED_AUTHOR_IDS)
            where_sql = 'author_id = %s'
            where_params = [author_id]
            if publish_clause:
                where_sql += ' AND ' + publish_clause
                where_params.extend(publish_params)
            if author_clause:
                where_sql += ' AND ' + author_clause
                where_params.extend(author_params)
            cursor.execute(f'SELECT * FROM video_info WHERE {where_sql}', tuple(where_params))
            rows = cursor.fetchall()
    finally:
        db_close(db)
    author_name = (rows[0].get('author_name') or '') if rows else ''
    return {
        'author_id': author_id,
        'author_name': author_name,
        'summary': analyzer.summarize_rows(rows),
        'trend': analyzer.build_trend(rows),
        'play_trend': analyzer.build_play_trend(rows),
        'top_videos': analyzer.top_videos(rows, sort_by=sort_by),
    }


@app.get('/api/cleanup/status')
def cleanup_status():
    cfg = _read_cleanup_config()
    return {
        'enabled': bool(cfg['enabled']),
        'last_clean_time': cfg['last_clean_time'],
        'batch_size': int(cfg['batch_size']),
        'authors': list(cfg['authors']),
    }


@app.post('/api/cleanup/toggle', dependencies=[Depends(verify_write_guard)])
def cleanup_toggle(req: CleanupToggleRequest):
    cfg = _read_cleanup_config()
    cfg['enabled'] = req.enabled
    _write_cleanup_config(cfg)
    return {'enabled': req.enabled}


@app.post('/api/cleanup/settings', dependencies=[Depends(verify_write_guard)])
def cleanup_settings(req: CleanupSettingsRequest):
    if not (1 <= req.batch_size <= 1000):
        raise HTTPException(status_code=400, detail='batch_size 必须在 1-1000 之间')
    cfg = _read_cleanup_config()
    cfg['batch_size'] = req.batch_size
    cfg['authors'] = list(req.authors)
    _write_cleanup_config(cfg)
    return {'batch_size': req.batch_size, 'authors': req.authors}


@app.get('/api/export')
def export_data(
    search: str = Query('', description='搜索视频标题/作者/ID'),
    sort_by: str = Query('crawl_time', description='排序字段'),
    order: str = Query('desc', pattern='^(asc|desc)$'),
    start_date: str = Query('', description='发布时间起始（YYYY-MM-DD）'),
    end_date: str = Query('', description='发布时间结束（YYYY-MM-DD）'),
    format: str = Query('csv', pattern='^(csv|xlsx)$'),
):
    allowed_sort = {
        'video_id', 'video_title', 'author_name', 'publish_time',
        'like_count', 'comment_count', 'share_count', 'play_count', 'collect_count',
        'crawl_time', 'update_time',
    }
    if sort_by not in allowed_sort:
        sort_by = 'crawl_time'
    order_clause = 'DESC' if order == 'desc' else 'ASC'
    publish_clause, publish_params = apply_publish_filter(start_date, end_date)
    where_parts = []
    params = []
    if search:
        where_parts.append('(video_id LIKE %s OR video_title LIKE %s OR author_name LIKE %s)')
        params.extend([f'%{search}%'] * 3)
    if publish_clause:
        where_parts.append(publish_clause)
        params.extend(publish_params)
    where_sql = ('WHERE ' + ' AND '.join(where_parts)) if where_parts else ''

    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute(f'SELECT COUNT(*) AS n FROM video_info {where_sql}', tuple(params))
            total = cursor.fetchone()['n']
        _check_export_total(total)
        if format == 'csv':
            def gen():
                conn = get_db()
                try:
                    with conn.cursor(pymysql.cursors.SSCursor) as cursor:
                        cursor.execute(
                            f'SELECT * FROM video_info {where_sql} ORDER BY {sort_by} {order_clause}',
                            tuple(params),
                        )
                        buf = io.StringIO()
                        buf.write('\ufeff')
                        writer = csv.DictWriter(buf, fieldnames=export_service.EXPORT_COLUMNS, extrasaction='ignore')
                        writer.writeheader()
                        yield buf.getvalue()
                        while True:
                            batch = cursor.fetchmany(1000)
                            if not batch:
                                break
                            buf = io.StringIO()
                            writer = csv.DictWriter(buf, fieldnames=export_service.EXPORT_COLUMNS, extrasaction='ignore')
                            for row in batch:
                                writer.writerow({c: ('' if row.get(c) is None else row.get(c)) for c in export_service.EXPORT_COLUMNS})
                            yield buf.getvalue()
                finally:
                    db_close(conn)
            return StreamingResponse(
                gen(), media_type='text/csv; charset=utf-8',
                headers={'Content-Disposition': 'attachment; filename="douyin_data.csv"'},
            )
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        from openpyxl import Workbook
        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        ws.append(list(export_service.EXPORT_COLUMNS))
        conn = get_db()
        try:
            with conn.cursor(pymysql.cursors.SSCursor) as cursor:
                cursor.execute(
                    f'SELECT * FROM video_info {where_sql} ORDER BY {sort_by} {order_clause}',
                    tuple(params),
                )
                while True:
                    batch = cursor.fetchmany(1000)
                    if not batch:
                        break
                    for row in batch:
                        ws.append([('' if row.get(c) is None else row.get(c)) for c in export_service.EXPORT_COLUMNS])
        finally:
            db_close(conn)
        wb.save(tmp.name)
        return FileResponse(
            tmp.name,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename='douyin_data.xlsx',
        )
    finally:
        db_close(db)


@app.post('/api/extension/videos', dependencies=[Depends(verify_write_guard)])
def extension_receive(req: ExtensionVideosRequest):
    valid, rejected = extension_receiver.validate_batch(req.model_dump())
    if not valid and not rejected:
        raise HTTPException(status_code=400, detail='没有可处理的记录')
    records = extension_receiver.dedupe_records(valid)
    records, author_rejected = extension_receiver.filter_by_author_whitelist(records, ALLOWED_AUTHOR_IDS)
    rejected.extend(author_rejected)
    db = get_db()
    try:
        with db.cursor() as cursor:
            for record in records:
                sql, params = extension_receiver.build_upsert(record)
                cursor.execute(sql, params)
            db.commit()
    finally:
        db_close(db)
    return {
        'source_url': req.source_url,
        'accepted': len(valid),
        'upserted': len(records),
        'rejected': rejected,
    }


FRONTEND_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'frontend')
DIST_DIR = os.path.join(FRONTEND_DIR, 'dist')

if os.path.isdir(DIST_DIR):
    app.mount('/app/assets', StaticFiles(directory=os.path.join(DIST_DIR, 'assets')), name='frontend-assets')

    @app.get('/app')
    @app.get('/app/{full_path:path}')
    def frontend_spa(full_path: str = ''):
        index = os.path.join(DIST_DIR, 'index.html')
        if full_path:
            target = os.path.abspath(os.path.join(DIST_DIR, full_path))
            if target.startswith(os.path.abspath(DIST_DIR)) and os.path.isfile(target):
                return FileResponse(target)
        return FileResponse(index)


@app.get('/', include_in_schema=False)
def root_redirect():
    return RedirectResponse('/app/')
